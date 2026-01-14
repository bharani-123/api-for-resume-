import csv
import os
import re
import json
import pandas as pd
import pdfplumber
import google.generativeai as genai
from docx import Document
from PIL import Image
import io
from flask import Flask, render_template, request, redirect, url_for, send_file, jsonify
from werkzeug.utils import secure_filename
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy.exc import IntegrityError
from datetime import date, datetime, timedelta
from urllib.parse import quote_plus
from flask import send_from_directory

# New Import for PDF Generation
try:
    from fpdf import FPDF
except ImportError:
    print("Warning: 'fpdf' module not found. Install it using 'pip install fpdf' to use PDF export.")
    FPDF = None

# ==========================================
# CONFIGURATION
# ==========================================
UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'png', 'jpg', 'jpeg', 'webp'}

initial_api_key = os.environ.get("GOOGLE_API_KEY")
if initial_api_key:
    genai.configure(api_key=initial_api_key)
    model = genai.GenerativeModel('gemini-2.5-flash')
else:
    model = None

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
# Increase upload limit to 100MB for bulk files
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024 

# ==========================================
# DATABASE CONFIGURATION
# ==========================================
database_url = os.environ.get('DATABASE_URL')
if not database_url:
    DB_USER = "root"
    DB_PASS = "2207" 
    DB_HOST = "localhost"
    DB_NAME = "resume_db"
    encoded_pass = quote_plus(DB_PASS)
    database_url = f"mysql+pymysql://{DB_USER}:{encoded_pass}@{DB_HOST}/{DB_NAME}"

app.config['SQLALCHEMY_DATABASE_URI'] = database_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'pool_recycle': 280, 'pool_pre_ping': True}

db = SQLAlchemy(app)
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# ==========================================
# DATABASE MODEL
# ==========================================
class Candidate(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(150))
    name = db.Column(db.String(100))
    email = db.Column(db.String(100), unique=True, index=True)
    phone = db.Column(db.String(50), unique=True, index=True)
    college = db.Column(db.String(200))
    degree = db.Column(db.String(100))
    department = db.Column(db.String(100))
    state = db.Column(db.String(50))
    district = db.Column(db.String(50))
    year_passing = db.Column(db.String(20))
    upload_date = db.Column(db.DateTime, default=datetime.utcnow)

# ==========================================
# HELPER FUNCTIONS
# ==========================================

def clean_list(data):
    """Converts ['Item A', 'Item B'] into 'Item A, Item B' to prevent database errors."""
    if isinstance(data, list):
        return ", ".join(str(item) for item in data)
    return data

def extract_text_traditional(file_path):
    ext = file_path.rsplit('.', 1)[1].lower()
    text = ""
    try:
        if ext == 'pdf':
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text += (page.extract_text() or "") + "\n"
        elif ext == 'docx':
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        print(f"Error reading document {file_path}: {e}")
    return text

def extract_name(text):
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    for line in lines[:10]:
        if re.match(r'^[A-Z][A-Z\s\.]{2,}$', line): return line.title()
    for line in lines[:10]:
        if re.match(r'^[A-Z]\.?( )?[A-Z][a-zA-Z]+$', line): return line.title()
    for line in lines[:10]:
        if re.match(r'^[A-Z][a-zA-Z]+ [A-Z][a-zA-Z]+$', line): return line.strip()
    return "Unknown"

def extract_email(text):
    m = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', text)
    return m.group(0) if m else "Not Specified"

def extract_phone(text):
    m = re.search(r'\b(?:\+?91)?\s*\d{10}\b', text)
    return m.group(0) if m else "Not Specified"

def extract_college(text):
    m = re.search(r'([A-Za-z ]+(University|Institute|College))', text, re.IGNORECASE)
    return m.group(0).strip() if m else "Not Specified"

def extract_degree(text):
    patterns = [r'b\.?tech', r'b\.?e', r'm\.?tech', r'm\.?e', r'bachelor(?: of)?', r'master(?: of)?',
            r'b\.?sc', r'm\.?sc', r'b\.?a', r'm\.?a', r'b\.?com', r'm\.?com', r'b\.?ba', r'm\.?ba',
            r'b\.?ca', r'm\.?ca', r'b\.?ed', r'm\.?ed', r'b\.?pharm', r'm\.?pharm', r'b\.?arch', r'm\.?arch',
            r'b\.?ds', r'm\.?ds', r'mbbs', r'bams', r'bhms', r'b\.?voc', r'm\.?voc', r'diploma', r'pg diploma',
            r'ph\.?d', r'doctorate']
    for p in patterns:
        match = re.search(p, text, re.IGNORECASE)
        if match: return match.group(0).upper()
    return "Not Specified"

def extract_department(text):
    txt_lower = text.lower()
    edu_section = text
    keywords = ["education", "academic details", "qualification", "educational qualification"]
    for key in keywords:
        if key in txt_lower:
            start = txt_lower.index(key)
            edu_section = text[start:start + 1000]
            break     
    patterns = [
        r"electronics and communication", r"ece", r"computer science", r"cs", r"cse",
        r"electrical and electronics", r"eee", r"mechanical engineering", r"mech", r"civil engineering", r"civil",
        r"artificial intelligence and data science", r"ai&ds", r"data science", r"data analytics",
        r"artificial intelligence", r"ai", r"cyber security", r"information technology", r"it",
        r"physics", r"chemistry", r"biology", r"biotechnology", r"mathematics", r"statistics",
        r"accounting", r"finance", r"banking", r"bba", r"marketing", r"hr",
        r"english", r"history", r"computer applications", r"ca", r"pharmacy", r"law"
    ]
    for p in patterns:
        if re.search(p, edu_section, re.IGNORECASE): return p.title()
    return "Not Specified"

def extract_state(text):
    states = ["Tamil Nadu", "Tamilnadu", "Kerala", "Karnataka", "Andhra Pradesh", "Telangana", "Maharashtra", "Delhi"]
    for state in states:
        if re.search(r'\b' + re.escape(state) + r'\b', text, re.IGNORECASE): return state.title()
    return "Not Specified"

def extract_district(text):
    districts = [r"chennai", r"coimbatore", r"madurai", r"trichy", r"salem", r"tirunelveli", r"erode", r"vellore", r"thoothukudi", r"dindigul", r"thanjavur", r"tiruppur", r"virudhunagar", r"karur", r"nilgiris", r"krishnagiri", r"kanyakumari", r"kancheepuram", r"namakkal", r"sivagangai", r"cuddalore", r"pudukkottai", r"theni", r"ramanathapuram", r"thiruvarur", r"thiruvallur", r"tiruvannamalai", r"nagapattinam", r"viluppuram", r"perambalur", r"dharmapuri", r"ariyalur", r"tirupathur", r"tenkasi", r"chengalpattu", r"kallakurichi", r"ranipet", r"mayiladuthurai"]
    for dist in districts:
        if re.search(r'\b' + re.escape(dist) + r'\b', text, re.IGNORECASE): return dist.title()
    return "Not Specified"

def extract_year_of_passing(text):
    pattern_range = r'(20\d{2})\s*[\-\–]\s*(\d{2,4})'
    match_range = re.search(pattern_range, text)
    if match_range:
        end_year = match_range.group(2)
        if len(end_year) == 2: return "20" + end_year 
        return end_year
    matches = re.findall(r'\b(20\d{2})\b', text)
    valid_years = [int(y) for y in matches if 2000 <= int(y) <= 2030]
    if valid_years: return str(max(valid_years))
    return "Not Specified"

def normalize_email(email):
    if not email: return None
    email = str(email).strip()
    if not email or email.lower() == 'not specified': return None
    return email.lower()

def normalize_phone(phone):
    if not phone: return None
    phone = str(phone).strip()
    if not phone or phone.lower() == 'not specified': return None
    digits = re.sub(r'\D', '', phone)
    if not digits: return None
    if len(digits) >= 10:
        return digits[-10:]
    return digits

def parse_with_regex(filepath):
    raw_text = extract_text_traditional(filepath)
    raw_text = raw_text.replace("■", "").replace("●", "")
    return {
        "Name": extract_name(raw_text),
        "Email": extract_email(raw_text),
        "Phone": extract_phone(raw_text),
        "College": extract_college(raw_text),
        "Degree": extract_degree(raw_text),
        "Department": extract_department(raw_text),
        "Passed Out": extract_year_of_passing(raw_text),
        "State": extract_state(raw_text),
        "District": extract_district(raw_text)
    }

def extract_data_with_gemini(file_path):
    if file_path.lower().endswith(('.png', '.jpg', '.jpeg', '.webp')):
        if model is None:
            return None 
        
        try:
            img = Image.open(file_path)
            prompt = """
             You are an expert Resume Parser. Analyze this resume image and extract the following details.
            Return ONLY a valid JSON object. Do not write markdown formatting.
            Fields: Name, Phone, Email, College, Degree, Department, District, State, Passed Out.
            If a value is not found, use "Not Specified".
            """
            response = model.generate_content([prompt, img])
            clean = response.text.strip().replace("```json", "").replace("```", "").strip()
            return json.loads(clean)
        except Exception as e:
            print("Error in Gemini:", e)
            return None
    return {"Name": "File format not supported"}


# ==========================================
# ROUTES
# ==========================================
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route('/')
def index():
    return render_template('index.html', extracted_data=None)

# ==========================================
# SERVE UPLOADED FILES
# ==========================================
@app.route('/uploads/<filename>')
def serve_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)


# ---------------------------------------------------------
# UPDATED UPLOAD ROUTE (One-by-one Processing)
# ---------------------------------------------------------
@app.route('/upload', methods=['POST'])
def upload_file():
    files_to_process = []
    
    # Check for Server Folder Upload
    folder_path = request.form.get('folder_path')
    if folder_path:
        folder_path = folder_path.strip()
        if os.path.isdir(folder_path):
            for root, dirs, filenames in os.walk(folder_path):
                for fname in filenames:
                    if allowed_file(fname):
                        files_to_process.append(os.path.join(root, fname))
    
    # Check for Direct File Upload
    if 'file' in request.files:
        files = request.files.getlist('file')
        for file in files:
            if file and allowed_file(file.filename):
                filename = secure_filename(file.filename)
                filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                file.save(filepath)
                files_to_process.append(filepath)

    if not files_to_process:
        return redirect(request.url)

    # Process files one by one to prevent 1 failure stopping all
    success_count = 0
    
    for filepath in files_to_process:
        try:
            filename = os.path.basename(filepath)
            ext = filename.rsplit('.', 1)[1].lower()
            data = None

            # 1. Extract Data
            if ext in ['pdf', 'docx']:
                data = parse_with_regex(filepath)
            elif ext in ['png', 'jpg', 'jpeg', 'webp']:
                data = extract_data_with_gemini(filepath)

            if not data:
                continue

            # 2. Normalize Keys
            name_val = data.get('Name') or 'Unknown'
            raw_email = data.get('Email')
            raw_phone = data.get('Phone')
            email_val = normalize_email(raw_email)
            phone_val = normalize_phone(raw_phone)

            # 3. Check for Duplicates (DATABASE QUERY FIRST)
            existing = None
            if email_val:
                existing = Candidate.query.filter_by(email=email_val).first()
            if not existing and phone_val:
                existing = Candidate.query.filter_by(phone=phone_val).first()

            # 4. Add or Update
            if existing:
                existing.filename = filename
                existing.name = name_val
                existing.email = email_val
                existing.phone = phone_val
                existing.college = clean_list(data.get('College'))
                existing.degree = clean_list(data.get('Degree'))
                existing.department = clean_list(data.get('Department'))
                existing.year_passing = clean_list(data.get('Passed Out'))
                existing.state = clean_list(data.get('State'))
                existing.district = clean_list(data.get('District'))
                existing.upload_date = datetime.utcnow()
                # No new Add needed, just commit update
            else:
                new_candidate = Candidate(
                    filename=filename,
                    name=name_val,
                    email=email_val,
                    phone=phone_val,
                    college=clean_list(data.get('College')), 
                    degree=clean_list(data.get('Degree')),
                    department=clean_list(data.get('Department')),
                    year_passing=clean_list(data.get('Passed Out')),
                    state=clean_list(data.get('State')),
                    district=clean_list(data.get('District'))
                )
                db.session.add(new_candidate)
            
            # 5. COMMIT IMMEDIATELY (Safety checkpoint)
            # This ensures if the NEXT file fails, this one is already saved.
            db.session.commit()
            success_count += 1
            print(f"Successfully processed: {filename}")

        except Exception as e:
            db.session.rollback() # Undo ONLY the failed file
            print(f"Failed to process {filename}: {e}")
            continue # Move to next file

    return redirect(url_for('dashboard'))

@app.route('/save_candidate', methods=['POST'])
def save_candidate():
    try:
        candidate_id = request.form.get('candidate_id')
        filename = request.form.get('filename')
        name = request.form.get('name')
        raw_email = request.form.get('email')
        raw_phone = request.form.get('phone')
        college = request.form.get('college')
        degree = request.form.get('degree')
        department = request.form.get('department')
        year_passing = request.form.get('year_passing')
        state = request.form.get('state')
        district = request.form.get('district')

        email = normalize_email(raw_email)
        phone = normalize_phone(raw_phone)

        candidate = None
        if candidate_id:
            candidate = Candidate.query.get(candidate_id)
        
        if not candidate:
            if email:
                candidate = Candidate.query.filter_by(email=email).first()
            if not candidate and phone:
                candidate = Candidate.query.filter_by(phone=phone).first()

        if candidate:
            candidate.filename = filename
            candidate.name = name
            candidate.email = email
            candidate.phone = phone
            candidate.college = college
            candidate.degree = degree
            candidate.department = department
            candidate.year_passing = year_passing
            candidate.state = state
            candidate.district = district
            candidate.upload_date = datetime.utcnow()
        else:
            candidate = Candidate(
                filename=filename, name=name, email=email, phone=phone,
                college=college, degree=degree, department=department,
                year_passing=year_passing, state=state, district=district
            )
            db.session.add(candidate)

        db.session.commit()
        return redirect(url_for('dashboard'))

    except Exception as e:
        db.session.rollback()
        return f"Database Error: {e}"

@app.route('/dashboard')
def dashboard():
    search_query = request.args.get('search')
    query = Candidate.query

    if search_query:
        search_term = f"%{search_query}%"
        query = query.filter(
            db.or_(
                Candidate.name.ilike(search_term),
                Candidate.email.ilike(search_term),
                Candidate.phone.ilike(search_term)
            )
        )

    candidates = query.order_by(Candidate.upload_date.desc()).all()
    return render_template('dashboard.html', candidates=candidates, search_query=search_query)

@app.route('/set_api_key', methods=['POST'])
def set_api_key():
    try:
        data = request.get_json() or {}
        key = data.get('api_key')
        if not key:
            return jsonify({'success': False, 'message': 'No API key provided'}), 400
        genai.configure(api_key=key)
        global model
        model = genai.GenerativeModel('gemini-2.5-flash')
        return jsonify({'success': True})
    except Exception as e:
        print('Error setting API key:', e)
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/edit/<int:id>')
def edit_candidate(id):
    candidate = Candidate.query.get_or_404(id)
    data = {
        'id': candidate.id,
        'filename': candidate.filename,
        'Name': candidate.name,
        'Email': candidate.email,
        'Phone': candidate.phone,
        'College': candidate.college,
        'Degree': candidate.degree,
        'Department': candidate.department,
        'Passed Out': candidate.year_passing,
        'State': candidate.state,
        'District': candidate.district
    }
    return render_template('index.html', extracted_data=data)

@app.route('/delete_bulk', methods=['POST'])
def delete_bulk():
    ids = request.form.getlist('selected_ids')
    if ids:
        try:
            Candidate.query.filter(Candidate.id.in_(ids)).delete(synchronize_session=False)
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            print("Error deleting:", e)
    return redirect(url_for('dashboard'))

def get_filtered_query():
    query = Candidate.query
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    if start_date and end_date:
        try:
            start = datetime.strptime(start_date, '%Y-%m-%d')
            end = datetime.strptime(end_date, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
            query = query.filter(Candidate.upload_date >= start, Candidate.upload_date <= end)
        except ValueError:
            pass 
            
    return query.all()

@app.route('/export/json')
def export_json():
    candidates = get_filtered_query()
    filename = "Resume_Data_Export.json"
    
    data = [{"Name": c.name, "Contact": c.phone, "Email": c.email, 
             "Degree": c.degree, "Department": c.department, "College": c.college, 
             "State": c.state, "District": c.district, "Passed Out": c.year_passing, 
             "File Name": c.filename} for c in candidates]
    
    json_str = json.dumps(data, indent=4)
    buf = io.BytesIO(json_str.encode('utf-8'))
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename, mimetype='application/json')

@app.route('/export/csv')
def export_csv():
    candidates = get_filtered_query()
    filename = "Resume_Data_Export.csv"
    
    data = [{"Name": c.name, "Contact": c.phone, "Email": c.email, 
             "Degree": c.degree, "Department": c.department, "College": c.college, 
             "State": c.state, "District": c.district, "Passed Out": c.year_passing, 
             "File Name": c.filename} for c in candidates]
    
    str_buf = io.StringIO()
    fieldnames = ["Name", "Contact", "Email", "Degree", "Department", "College", "State", "District", "Passed Out", "File Name"]
    writer = csv.DictWriter(str_buf, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(data)
    bytes_buf = io.BytesIO(str_buf.getvalue().encode('utf-8'))
    bytes_buf.seek(0)
    return send_file(bytes_buf, as_attachment=True, download_name=filename, mimetype='text/csv')

@app.route('/export/excel')
def export_excel():
    candidates = get_filtered_query()
    filename = "Resume_Data_Export.xlsx"
    
    data = [{"Name": c.name, "Contact": c.phone, "Email": c.email, 
             "Degree": c.degree, "Department": c.department, "College": c.college, 
             "State": c.state, "District": c.district, "Passed Out": c.year_passing, 
             "File Name": c.filename} for c in candidates]
    
    df = pd.DataFrame(data)
    bytes_buf = io.BytesIO()
    with pd.ExcelWriter(bytes_buf, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    bytes_buf.seek(0)
    return send_file(bytes_buf, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/export/pdf')
def export_pdf():
    if FPDF is None:
        return "FPDF library not installed. Please run 'pip install fpdf'", 500

    candidates = get_filtered_query()
    
    class PDF(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 14)
            self.cell(0, 10, 'Candidate Database Report', 0, 1, 'C')
            self.ln(5)
            self.set_fill_color(240, 240, 240)
            self.set_font('Arial', 'B', 10)
            self.cell(45, 10, 'Name', 1, 0, 'C', 1)
            self.cell(35, 10, 'Phone', 1, 0, 'C', 1)
            self.cell(65, 10, 'Email', 1, 0, 'C', 1)
            self.cell(45, 10, 'Degree', 1, 1, 'C', 1)

        def footer(self):
            self.set_y(-15)
            self.set_font('Arial', 'I', 8)
            self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=9)

    for c in candidates:
        name = (c.name or "N/A")[:22]
        phone = (c.phone or "N/A")[:15]
        email = (c.email or "N/A")[:35]
        degree = (c.degree or "N/A")[:22]
        
        pdf.cell(45, 10, name, 1)
        pdf.cell(35, 10, phone, 1)
        pdf.cell(65, 10, email, 1)
        pdf.cell(45, 10, degree, 1)
        pdf.ln()

    try:
        pdf_content = pdf.output(dest='S').encode('latin-1') 
    except:
        pdf_content = pdf.output(dest='S').encode('latin-1', errors='ignore')

    buffer = io.BytesIO(pdf_content)
    buffer.seek(0)
    
    return send_file(buffer, as_attachment=True, download_name='Resume_Report.pdf', mimetype='application/pdf')

@app.route('/init-db')
def init_db():
    with app.app_context():
        db.create_all()
    return "Database initialized!"

if __name__ == '__main__':
    app.run(debug=True)